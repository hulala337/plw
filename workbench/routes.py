from __future__ import annotations

from . import runtime
from .errors import install_handlers

install_handlers(runtime.api, runtime.logger)

runtime.api.mount(
    "/assets", runtime.StaticFiles(directory=str(runtime.WEB / "assets")), name="assets"
)


@runtime.api.get("/api/display-info")
def api_display_info():
    from . import desktop

    return desktop.display_info()


@runtime.api.get("/api/health")
def api_health():
    from . import database, desktop, tracking

    listener_ok = tracking.listeners_healthy()
    tracker_ok = any(
        (t.name == "tracker" and t.is_alive() for t in runtime.threading.enumerate())
    )
    db_ok = False
    db_error = ""
    try:
        conn = database.db()
        conn.execute("SELECT 1").fetchone()
        conn.close()
        db_ok = True
    except Exception as exc:
        db_error = str(exc)
    web_ok = runtime.WEB.is_dir() and (runtime.WEB / "index.html").is_file()
    checks = {
        "keyboard_listener": bool(
            runtime.listener_refs
            and len(runtime.listener_refs) >= 1
            and runtime.listener_refs[0].is_alive()
        ),
        "mouse_listener": bool(
            runtime.listener_refs
            and len(runtime.listener_refs) >= 2
            and runtime.listener_refs[1].is_alive()
        ),
        "tracker": tracker_ok,
        "database": db_ok,
        "web": web_ok,
        "display": isinstance(desktop.display_info().get("count"), int),
    }
    return {
        "status": "ok" if all(checks.values()) else "degraded",
        "listeners": listener_ok,
        **checks,
        "db_error": db_error,
        "last_listener_ok": runtime.listener_last_ok,
        "last_keyboard_event": runtime.listener_last_keyboard_event,
        "last_mouse_event": runtime.listener_last_mouse_event,
        "listener_restarts": runtime.listener_restart_count,
        "listener_error": runtime.listener_error,
    }


@runtime.api.get("/")
def index():
    return runtime.FileResponse(runtime.WEB / "index.html")


@runtime.api.get("/api/growth")
def growth():
    from . import services

    return services.growth_payload()


@runtime.api.patch("/api/equipment")
def patch_equipment(item: runtime.EquipmentPatch):
    from . import database

    slot = item.slot.strip().lower()
    allowed = {
        "scene": "scenes",
        "pelican": "pelicans",
        "outfit": "outfits",
        "accessory": "accessories",
        "decoration": "decorations",
        "effect": "effects",
    }
    if slot not in allowed:
        raise runtime.HTTPException(400, "不支持的装备槽位")
    conn = database.db()
    try:
        row = conn.execute(
            "SELECT id,required_level FROM " + allowed[slot] + " WHERE id=?",
            (item.item_id,),
        ).fetchone()
        if not row:
            raise runtime.HTTPException(404, "内容不存在")
        profile = conn.execute(
            "SELECT level FROM progress_profile WHERE id=1"
        ).fetchone()
        level = int(profile["level"] if profile else 1)
        unlocked = conn.execute(
            "SELECT 1 FROM unlocks WHERE item_type=? AND item_id=?",
            (slot, item.item_id),
        ).fetchone()
        if not unlocked and int(row["required_level"] or 1) > 1:
            raise runtime.HTTPException(403, "内容尚未解锁")
        if slot in {"pelican", "outfit"}:
            current_pelican = conn.execute(
                "SELECT item_id FROM equipment WHERE slot='pelican'"
            ).fetchone()
            target_pelican = (
                item.item_id
                if slot == "pelican"
                else current_pelican["item_id"]
                if current_pelican
                else "classic"
            )
            outfit = conn.execute(
                "SELECT pelican_id FROM outfits WHERE id=(SELECT item_id FROM equipment WHERE slot='outfit')"
            ).fetchone()
            target_outfit = (
                conn.execute(
                    "SELECT pelican_id FROM outfits WHERE id=?", (item.item_id,)
                ).fetchone()
                if slot == "outfit"
                else outfit
            )
            if (
                target_outfit
                and target_outfit["pelican_id"]
                and (target_outfit["pelican_id"] != target_pelican)
            ):
                raise runtime.HTTPException(409, "该鹈鹕与当前服装不匹配")
        now = runtime.datetime.now().isoformat(timespec="seconds")
        storage_slot = "decoration:" + item.item_id if slot == "decoration" else slot
        conn.execute(
            "INSERT INTO equipment(slot,item_type,item_id,updated_at) VALUES(?,?,?,?) ON CONFLICT(slot) DO UPDATE SET item_type=excluded.item_type,item_id=excluded.item_id,updated_at=excluded.updated_at",
            (storage_slot, slot, item.item_id, now),
        )
        conn.commit()
        equipment = [
            dict(x)
            for x in conn.execute("SELECT * FROM equipment ORDER BY slot").fetchall()
        ]
        return {"ok": True, "equipment": equipment}
    finally:
        conn.close()


@runtime.api.get("/api/world")
def world():
    from . import services

    return services.world_payload()


@runtime.api.get("/api/dashboard")
def dashboard(range: str = "today"):
    from . import services

    if range not in {"today", "week", "month"}:
        raise runtime.HTTPException(400, "invalid range")
    return services.api_payload(range)


@runtime.api.get("/api/settings")
def get_settings():
    from . import database, desktop

    return {
        "version": runtime.VERSION,
        "autostart": desktop.startup_enabled(),
        "idle_seconds": int(
            database.setting_get("idle_seconds", str(runtime.IDLE_SECONDS))
        ),
        "weather_enabled": False,
        "desktop_pet": database.setting_get("desktop_pet", "1") != "0",
        "sounds_enabled": database.setting_get("sounds_enabled", "0") != "0",
        "privacy": {
            "local_only": not database.setting_get("weather_enabled", "1") != "0",
            "actual_input": False,
            "window_titles": False,
            "urls": False,
            "weather_external": database.setting_get("weather_enabled", "1") != "0",
            "weather_provider": "Open-Meteo",
        },
    }


@runtime.api.patch("/api/settings")
def patch_settings(item: runtime.SettingsPatch):
    from . import database, desktop, routes

    if item.autostart is not None:
        desktop.set_startup(bool(item.autostart))
    if item.idle_seconds is not None:
        database.setting_set(
            "idle_seconds", str(max(30, min(900, int(item.idle_seconds))))
        )
    if item.weather_enabled is not None:
        database.setting_set("weather_enabled", "1" if item.weather_enabled else "0")
    if item.desktop_pet is not None:
        database.setting_set("desktop_pet", "1" if item.desktop_pet else "0")
    if item.sounds_enabled is not None:
        database.setting_set("sounds_enabled", "1" if item.sounds_enabled else "0")
    return routes.get_settings()


@runtime.api.post("/api/todos")
def add_todo(item: runtime.TodoIn):
    from . import database

    title = item.title.strip()
    if not title:
        raise runtime.HTTPException(400, "任务不能为空")
    conn = database.db()
    cur = conn.execute(
        "INSERT INTO todos(title,created_at) VALUES(?,?)",
        (title, runtime.datetime.now().isoformat(timespec="seconds")),
    )
    conn.commit()
    tid = cur.lastrowid
    conn.close()
    return {"id": tid, "title": title, "done": 0}


@runtime.api.patch("/api/todos/{todo_id}")
def toggle_todo(todo_id: int, item: runtime.TodoPatch):
    from . import database

    conn = database.db()
    try:
        row = conn.execute(
            "SELECT done,completion_count FROM todos WHERE id=?", (todo_id,)
        ).fetchone()
        if not row:
            raise runtime.HTTPException(404, "任务不存在")
        was_done = bool(row["done"])
        now = runtime.datetime.now().isoformat(timespec="microseconds")
        if item.done and (not was_done):
            count = int(row["completion_count"] or 0) + 1
            conn.execute(
                "UPDATE todos SET done=1,completed_at=?,completion_count=? WHERE id=?",
                (now, count, todo_id),
            )
        elif not item.done:
            conn.execute(
                "UPDATE todos SET done=0,completed_at=NULL WHERE id=?", (todo_id,)
            )
        conn.commit()
        return {
            "ok": True,
            "done": bool(item.done),
            "completion_count": int(row["completion_count"] or 0)
            + (1 if item.done and (not was_done) else 0),
        }
    finally:
        conn.close()


@runtime.api.delete("/api/todos/{todo_id}")
def delete_todo(todo_id: int):
    from . import database

    conn = database.db()
    cur = conn.execute("DELETE FROM todos WHERE id=?", (todo_id,))
    conn.commit()
    conn.close()
    if not cur.rowcount:
        raise runtime.HTTPException(404, "任务不存在")
    return {"ok": True}


@runtime.api.post("/api/forget-today")
def forget_today():
    from . import database, services

    target = services.today_key()
    start = target + "T00:00:00"
    end = (runtime.date.today() + runtime.timedelta(days=1)).isoformat() + "T00:00:00"
    with runtime.pending_lock:
        for key in runtime.pending:
            runtime.pending[key] = 0
        runtime.pending_time["active_seconds"] = 0.0
        runtime.pending_time["idle_seconds"] = 0.0
    runtime.last_input_ts = 0.0
    runtime.last_xy = None
    runtime.last_monitor_index = None
    runtime.tracker_reset_seq += 1
    conn = database.db()
    from . import key_counts
    with key_counts.lock:
        for entry in list(key_counts.pending):
            if entry[0] == target:
                del key_counts.pending[entry]
        conn.execute("DELETE FROM daily_key_counts WHERE day=?", (target,))
    conn.execute("DELETE FROM character_receipts WHERE day=?", (target,))
    conn.execute("DELETE FROM daily WHERE day=?", (target,))
    conn.execute("DELETE FROM activity_slices WHERE date(slice_start)=?", (target,))
    conn.execute("DELETE FROM app_usage WHERE day=?", (target,))
    conn.execute(
        "DELETE FROM focus_sessions WHERE started_at < ? AND (ended_at IS NULL OR ended_at >= ?)",
        (end, start),
    )
    conn.commit()
    conn.close()
    database.ensure_today(target)
    return {"ok": True, "day": target}


@runtime.api.get("/api/replay")
def replay(day: str | None = None):
    from . import database, services

    if day is not None:
        try:
            parsed = runtime.date.fromisoformat(day)
            if parsed.isoformat() != day:
                raise ValueError("Expected YYYY-MM-DD")
        except ValueError:
            raise runtime.HTTPException(400, "invalid day; expected YYYY-MM-DD")
    target = day or services.today_key()
    conn = database.db()
    rows = conn.execute(
        "SELECT slice_start,category,seconds,events FROM activity_slices WHERE date(slice_start)=? ORDER BY slice_start",
        (target,),
    ).fetchall()
    conn.close()
    return {"day": target, "events": [dict(r) for r in rows]}


@runtime.api.get("/api/export/csv")
def export_csv(range: str = "today"):
    from . import database, services

    if range not in {"today", "week", "month"}:
        raise runtime.HTTPException(400, "invalid range")
    start, end = services.range_bounds(
        range if range in {"today", "week", "month"} else "today"
    )
    daily, tl, apps, sessions = database.export_rows(start, end)
    out = runtime.io.StringIO()
    w = runtime.csv.writer(out)
    w.writerow(
        [
            "类型",
            "日期/时间",
            "分类",
            "秒数",
            "键盘",
            "字符数",
            "左键",
            "右键",
            "中键",
            "滚轮",
            "光标像素",
            "活跃秒",
            "空闲秒",
            "行为事件",
            "跨屏次数",
            "会话分类",
            "结束原因",
        ]
    )
    for r in daily:
        w.writerow(
            [
                "daily",
                r["day"],
                "",
                "",
                r["keys"],
                r["text_chars"],
                r["left_click"],
                r["right_click"],
                r["middle_click"],
                r["scroll_events"],
                round(r["cursor_distance_px"]),
                round(r["active_seconds"]),
                round(r["idle_seconds"]),
                r["activity_events"],
                r["monitor_switches"],
                "",
                "",
            ]
        )
    for r in tl:
        w.writerow(
            [
                "timeline",
                r["slice_start"],
                runtime.CATEGORY_LABELS.get(r["category"], r["category"]),
                r["seconds"],
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )
    for r in apps:
        w.writerow(
            [
                "usage",
                r["day"],
                runtime.CATEGORY_LABELS.get(r["category"], r["category"]),
                round(r["seconds"]),
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )
    for r in sessions:
        w.writerow(
            [
                "focus_session",
                r["started_at"],
                "",
                round(r["active_seconds"]),
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                r["categories"],
                r["ended_reason"],
            ]
        )
    data = out.getvalue().encode("utf-8-sig")
    return runtime.StreamingResponse(
        runtime.io.BytesIO(data),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="pelican_{range}.csv"'},
    )


@runtime.api.get("/api/export/xlsx")
def export_xlsx(range: str = "today"):
    from . import database, services

    if range not in {"today", "week", "month"}:
        raise runtime.HTTPException(400, "invalid range")
    from openpyxl import Workbook

    start, end = services.range_bounds(
        range if range in {"today", "week", "month"} else "today"
    )
    daily, tl, apps, sessions = database.export_rows(start, end)
    wb = Workbook()
    ws = wb.active
    ws.title = "日报"
    ws.append(
        [
            "日期",
            "键盘",
            "输入字符数",
            "鼠标左键",
            "鼠标右键",
            "滚轮",
            "鼠标移动(px)",
            "活跃(s)",
            "空闲(s)",
            "行为事件",
            "跨屏次数",
        ]
    )
    for r in daily:
        ws.append(
            [
                r["day"],
                r["keys"],
                r["text_chars"],
                r["left_click"],
                r["right_click"],
                r["scroll_events"],
                round(r["cursor_distance_px"]),
                round(r["active_seconds"]),
                round(r["idle_seconds"]),
                r["activity_events"],
                r["monitor_switches"],
            ]
        )
    ws2 = wb.create_sheet("时间轴")
    ws2.append(["时间", "分类", "秒数"])
    for r in tl:
        ws2.append(
            [
                r["slice_start"],
                runtime.CATEGORY_LABELS.get(r["category"], r["category"]),
                r["seconds"],
            ]
        )
    ws3 = wb.create_sheet("专注Session")
    ws3.append(["开始", "结束", "活跃秒", "分类序列", "结束原因"])
    for r in sessions:
        ws3.append(
            [
                r["started_at"],
                r["ended_at"],
                r["active_seconds"],
                r["categories"],
                r["ended_reason"],
            ]
        )
    ws4 = wb.create_sheet("工作类型")
    ws4.append(["日期", "分类", "秒数"])
    for r in apps:
        ws4.append(
            [
                r["day"],
                runtime.CATEGORY_LABELS.get(r["category"], r["category"]),
                round(r["seconds"]),
            ]
        )
    mem = runtime.io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return runtime.StreamingResponse(
        mem,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="pelican_{range}.xlsx"'},
    )


from .characters import CharacterReceipt

@runtime.api.post("/api/characters")
def record_characters(item: CharacterReceipt):
    from . import characters, database
    try:
        return characters.receive(database.db, item)
    except ValueError as exc:
        raise runtime.HTTPException(400, str(exc))
